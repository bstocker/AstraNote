"""Facturation des séances et bilan NDA."""
from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from astranote.models import (
    db, School, AcademicYear, Class, Module, Student, Enrollment, GradeDate,
)
from conftest import make_teacher, login


def build(app, teacher_id=1):
    """École ESIEE-IT, deux classes de même niveau (L3A : 3 inscrits dont un
    neutralisé, L3B : 2 inscrits, dont un aussi inscrit en L3A)."""
    with app.app_context():
        school = School(name="ESIEE-IT", billing_emails="compta@esiee.fr",
                        billing_cc_emails="dir@esiee.fr", contract_number="CT-42")
        year = AcademicYear(label="2025-2026")
        db.session.add_all([school, year])
        db.session.flush()
        l3a = Class(name="L3A", school_id=school.id, academic_year_id=year.id,
                    teacher_id=teacher_id, level="Bachelor", hourly_rate=50)
        l3b = Class(name="L3B", school_id=school.id, academic_year_id=year.id,
                    teacher_id=teacher_id, level="Bachelor")
        db.session.add_all([l3a, l3b])
        db.session.flush()
        alice, bob, zoe = Student(full_name="Alice"), Student(full_name="Bob"), \
            Student(full_name="Zoé", active=False)
        carl = Student(full_name="Carl")
        db.session.add_all([alice, bob, zoe, carl])
        db.session.flush()
        for klass, students in ((l3a, (alice, bob, zoe)), (l3b, (alice, carl))):
            for st in students:
                db.session.add(Enrollment(student_id=st.id, class_id=klass.id))
        ids = {}
        for klass in (l3a, l3b):
            mp = Module(name="Mngt projet", class_id=klass.id)
            suivi = Module(name="Suivi", class_id=klass.id)
            db.session.add_all([mp, suivi])
            db.session.flush()
            # 21 h de cours (7 + 7 + 7) et 15 h de suivi, en 2026.
            for i, d in enumerate((date(2026, 1, 5), date(2026, 1, 12), date(2026, 1, 19))):
                db.session.add(GradeDate(module_id=mp.id, date=d, duration_hours=7,
                                         label=f"S{i + 1}", position=i))
            db.session.add(GradeDate(module_id=suivi.id, date=date(2026, 2, 2),
                                     duration_hours=15))
            ids[klass.name] = {"mp": mp.id, "suivi": suivi.id, "class": klass.id}
        db.session.commit()
        ids["year"] = year.id
        return ids


def rows_of(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


# --------------------------------------------------------------------------- #
# Page et saisie
# --------------------------------------------------------------------------- #
def test_billing_page_lists_done_sessions_only(app, admin):
    ids = build(app)
    with app.app_context():
        db.session.add(GradeDate(module_id=ids["L3A"]["mp"], label="Future",
                                 date=date.today() + timedelta(days=30), duration_hours=3))
        db.session.add(GradeDate(module_id=ids["L3A"]["mp"], label="SansDate"))
        db.session.commit()
    html = admin.get(f"/facturation?year={ids['year']}").get_data(as_text=True)
    assert "Mngt projet" in html and "compta@esiee.fr" in html and "CT-42" in html
    assert "SansDate" in html
    assert "Future" not in html
    assert "21 h réalisées" in html


def test_save_session_hours_and_ref(app, admin):
    ids = build(app)
    with app.app_context():
        gd = GradeDate.query.filter_by(module_id=ids["L3A"]["mp"]).first()
        gid = gd.id
    r = admin.post(f"/facturation/seances/{gid}", json={"duration_hours": "3,5"})
    assert r.status_code == 200 and r.get_json()["billed"] is False
    r = admin.post(f"/facturation/seances/{gid}", json={"billing_ref": "  F-2026-01 "})
    assert r.get_json()["billed"] is True
    with app.app_context():
        gd = db.session.get(GradeDate, gid)
        assert gd.duration_hours == 3.5 and gd.billing_ref == "F-2026-01"

    # La séance facturée est grisée au rechargement, et jamais cochée.
    html = admin.get(f"/facturation?year={ids['year']}").get_data(as_text=True)
    assert 'class="billing-row billed"' in html
    assert "checked" not in html.split('id="billingForm"')[1].split("Bilan NDA")[0]

    # Référence effacée : la séance redevient à facturer.
    r = admin.post(f"/facturation/seances/{gid}", json={"billing_ref": "   "})
    assert r.get_json()["billed"] is False
    assert admin.post(f"/facturation/seances/{gid}",
                      json={"duration_hours": "99"}).status_code == 400


def test_billing_scoped_to_teacher(app, admin):
    ids = build(app)
    make_teacher(app, "Autre", "autre@x.fr")
    client = app.test_client()
    login(client, "autre@x.fr")
    with app.app_context():
        gid = GradeDate.query.first().id
    assert client.post(f"/facturation/seances/{gid}",
                       json={"billing_ref": "x"}).status_code == 403
    assert client.get(f"/facturation/export.xlsx?sessions={gid}").status_code == 403
    assert "Mngt projet" not in client.get("/facturation").get_data(as_text=True)


# --------------------------------------------------------------------------- #
# Extraction du lot
# --------------------------------------------------------------------------- #
def test_export_billing_batch(app, admin):
    ids = build(app)
    with app.app_context():
        l3a = GradeDate.query.filter_by(module_id=ids["L3A"]["mp"]).all()
        suivi_b = GradeDate.query.filter_by(module_id=ids["L3B"]["suivi"]).one()
        l3a[0].billing_ref = "F-01"
        db.session.commit()
        chosen = [l3a[0].id, l3a[1].id, suivi_b.id]
    qs = "&".join(f"sessions={i}" for i in chosen)
    r = admin.get(f"/facturation/export.xlsx?{qs}")
    assert r.status_code == 200
    rows = rows_of(load_workbook(BytesIO(r.data)).active)
    flat = [c for row in rows for c in row if c is not None]
    assert rows[2][0] == "Total des heures" and rows[2][2] == 29   # 7 + 7 + 15
    assert "ESIEE-IT · L3A · Mngt projet" in flat
    assert "ESIEE-IT · L3B · Suivi" in flat
    assert "compta@esiee.fr" in flat and "dir@esiee.fr" in flat and "CT-42" in flat
    assert "F-01" in flat
    subtotals = [row[2] for row in rows if row[1] == "Sous-total"]
    assert subtotals == [14, 15]
    # Montant uniquement pour la classe qui a un taux horaire (L3A, 50 €/h).
    assert [row[2] for row in rows if row[1] == "Montant (€)"] == [700]


def test_export_billing_without_selection_redirects(app, admin):
    build(app)
    assert admin.get("/facturation/export.xlsx").status_code == 302


# --------------------------------------------------------------------------- #
# Bilan NDA
# --------------------------------------------------------------------------- #
def test_nda_report_matches_trainee_hours_rule(app, admin):
    """(21 h + 15 h) × effectif de chaque classe, cumulé par école et niveau."""
    ids = build(app)
    r = admin.get("/facturation/nda.xlsx?start=2026-01-01&end=2026-12-31")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.data))

    synth = {row[0]: row[1] for row in rows_of(wb["Synthèse"]) if row[0]}
    # Alice (2 classes) comptée une fois, Zoé neutralisée exclue : Alice, Bob, Carl.
    assert synth["Cadre G — Étudiants uniques (inscrits, non neutralisés)"] == 3
    assert synth["Heures d'enseignement (séances de la période)"] == 72   # 36 h × 2 classes
    # L3A : 36 × 2 actifs = 72 ; L3B : 36 × 2 = 72.
    assert synth["Cadre G — Heures-stagiaires (heures × étudiants)"] == 144
    # CA : seule L3A a un taux (36 h × 50 €).
    assert synth["Cadre C — CA estimé (heures × taux horaire, €)"] == 1800
    assert any("sans taux horaire" in str(k) for k in synth)

    group = rows_of(wb["École & niveau"])
    assert group[1][:6] == ["ESIEE-IT", "Bachelor", 3, 72, 144, 1800]

    modules = rows_of(wb["Par module"])
    mp_a = next(r for r in modules if r[2] == "L3A" and r[3] == "Mngt projet")
    assert mp_a[5] == 21 and mp_a[7] == 2 and mp_a[8] == 42 and mp_a[10] == 1050


def test_nda_period_filters_sessions(app, admin):
    build(app)
    r = admin.get("/facturation/nda.xlsx?start=2026-01-01&end=2026-01-31")
    synth = {row[0]: row[1] for row in rows_of(load_workbook(BytesIO(r.data))["Synthèse"])
             if row[0]}
    # Janvier : les 21 h de cours seulement, pas le suivi de février.
    assert synth["Heures d'enseignement (séances de la période)"] == 42


def test_class_level_editable(app, admin):
    ids = build(app)
    cid = ids["L3A"]["class"]
    admin.post(f"/classes/{cid}/billing", data={"hourly_rate": "50", "level": "Mastère"})
    with app.app_context():
        assert db.session.get(Class, cid).level == "Mastère"
    # Formulaire sans champ niveau : niveau conservé.
    admin.post(f"/classes/{cid}/billing", data={"hourly_rate": "55"})
    with app.app_context():
        assert db.session.get(Class, cid).level == "Mastère"
