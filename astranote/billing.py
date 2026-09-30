"""Facturation des séances et bilan NDA (bilan pédagogique et financier).

Deux usages, une même matière première — les séances et leur durée :

- **Facturation** : les séances réalisées des modules d'une année, module par
  module. On en coche un lot, on l'extrait en Excel avec le total des heures,
  et l'on note sur chaque séance la référence de la facture (n° libre). Une
  séance qui porte une référence est considérée comme facturée : grisée, et
  exclue du « tout cocher » pour ne pas brouiller le total de ce qui reste à
  facturer. La coche n'est jamais enregistrée : elle ne vaut que pour le lot
  en cours.
- **Bilan NDA** : effectifs et heures-stagiaires d'une période (cadre G), et
  chiffre d'affaires estimé à partir du taux horaire des classes (cadre C).
"""
from datetime import date as date_type, datetime
from io import BytesIO

from flask import (
    Blueprint, render_template, request, jsonify, abort, send_file, flash,
    redirect, url_for,
)
from flask_login import login_required

from .models import db, Class, GradeDate
from .main import (
    visible_classes_query, visible_years_query, selected_year_id,
    get_class_or_403,
)
from .modules import module_dates_sorted, parse_duration, MAX_DURATION_HOURS

billing_bp = Blueprint("billing", __name__)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NO_LEVEL = "Niveau non précisé"


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def is_billed(gd):
    """Une séance est facturée dès qu'une référence non vide lui est associée."""
    return bool((gd.billing_ref or "").strip())


def done_sessions(module, today=None):
    """Séances réalisées d'un module, de la plus ancienne à la plus récente.

    Une séance à venir n'a rien à facturer. Une séance sans date est gardée :
    rien ne dit qu'elle n'a pas eu lieu, et l'écarter la ferait oublier.
    """
    today = today or date_type.today()
    return [gd for gd in module_dates_sorted(module)
            if gd.date is None or gd.date <= today]


def _sorted_classes(classes):
    return sorted(classes, key=lambda c: (c.school.name.lower(), c.name.lower()))


def _sum_hours(sessions):
    return round(sum(gd.duration_hours or 0 for gd in sessions), 2)


def _money(value):
    return round(value, 2) if value is not None else None


def _parse_day(raw, default):
    try:
        return datetime.strptime((raw or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        return default


def _xlsx_response(wb, filename):
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=filename,
                     mimetype=XLSX_MIME)


# --------------------------------------------------------------------------- #
# Page de facturation
# --------------------------------------------------------------------------- #
@billing_bp.route("/facturation")
@login_required
def view_billing():
    years = visible_years_query().all()
    year_id = selected_year_id(years)

    blocks = []
    if year_id:
        classes = visible_classes_query().filter(
            Class.academic_year_id == year_id).all()
        for klass in _sorted_classes(classes):
            for module in sorted(klass.modules, key=lambda m: m.name.lower()):
                sessions = done_sessions(module)
                if not sessions:
                    continue
                blocks.append({
                    "module": module, "klass": klass, "school": klass.school,
                    "sessions": sessions,
                    "hours": _sum_hours(sessions),
                    "unbilled": _sum_hours([s for s in sessions if not is_billed(s)]),
                })

    # Bilan NDA : l'exercice civil en cours par défaut.
    this_year = date_type.today().year
    return render_template(
        "billing/billing.html", years=years, selected_year_id=year_id,
        blocks=blocks, is_billed=is_billed,
        nda_start=date_type(this_year, 1, 1), nda_end=date_type(this_year, 12, 31),
    )


@billing_bp.route("/facturation/seances/<int:date_id>", methods=["POST"])
@login_required
def save_session(date_id):
    """Durée et/ou référence de facturation d'une séance (saisie AJAX).

    La durée est celle de la séance elle-même : la corriger ici la corrige
    aussi dans la grille du module et dans le cumul d'heures.
    """
    gd = db.session.get(GradeDate, date_id) or abort(404)
    get_class_or_403(gd.module.class_id)
    data = request.get_json(silent=True) or {}

    if "duration_hours" in data:
        try:
            gd.duration_hours = parse_duration(str(data["duration_hours"] or ""))
        except ValueError:
            return jsonify(error=f"Durée invalide : indiquez un nombre d'heures "
                                 f"entre 0 et {MAX_DURATION_HOURS}."), 400
    if "billing_ref" in data:
        ref = str(data["billing_ref"] or "").strip()
        gd.billing_ref = ref[:120] or None
    db.session.commit()
    return jsonify(ok=True, duration_hours=gd.duration_hours,
                   billing_ref=gd.billing_ref or "", billed=is_billed(gd))


@billing_bp.route("/facturation/export.xlsx")
@login_required
def export_billing():
    """Lot de séances cochées : module par module (école, contacts, séances,
    texte), sous-total d'heures par module et total général."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    ids = {int(i) for i in request.args.getlist("sessions") if str(i).isdigit()}
    sessions = GradeDate.query.filter(GradeDate.id.in_(ids or [0])).all()
    # Périmètre : une séance d'une classe d'un autre enseignant est refusée
    # plutôt qu'ignorée — elle ne peut venir que d'une URL fabriquée.
    for gd in sessions:
        get_class_or_403(gd.module.class_id)
    if not sessions:
        flash("Aucune séance cochée : rien à extraire.", "error")
        return redirect(url_for("billing.view_billing",
                                year=request.args.get("year", type=int)))

    by_module = {}
    for gd in sessions:
        by_module.setdefault(gd.module, []).append(gd)
    modules = sorted(by_module, key=lambda m: (
        m.klass.school.name.lower(), m.klass.name.lower(), m.name.lower()))

    wb = Workbook()
    ws = wb.active
    ws.title = "Facturation"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="F1F5F9")
    total_fill = PatternFill("solid", fgColor="FEF9C3")
    right = Alignment(horizontal="right")

    total_hours = _sum_hours(sessions)
    ws["A1"] = f"Facturation — lot du {date_type.today().strftime('%d/%m/%Y')}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"{len(sessions)} séance(s) · {len(modules)} module(s)"
    ws["A3"] = "Total des heures"
    ws["A3"].font = bold
    ws["C3"] = total_hours
    ws["C3"].font = bold
    ws["C3"].fill = total_fill
    row = 5

    for module in modules:
        klass, school = module.klass, module.klass.school
        rows = sorted(by_module[module],
                      key=lambda g: (g.date is None, g.date or date_type.min,
                                     g.position or 0))
        ws.cell(row=row, column=1,
                value=f"{school.name} · {klass.name} · {module.name}").font = Font(bold=True, size=12)
        row += 1
        for label, value in (
            ("Contacts de facturation", school.billing_emails),
            ("Contacts en copie (CC)", school.billing_cc_emails),
            ("N° de contrat", school.contract_number),
            ("Taux horaire (€/h)", klass.hourly_rate),
        ):
            if value not in (None, ""):
                ws.cell(row=row, column=1, value=label)
                ws.cell(row=row, column=2, value=value)
                row += 1
        for j, title in enumerate(("Date", "Libellé", "Heures", "N° / texte"), start=1):
            c = ws.cell(row=row, column=j, value=title)
            c.font, c.fill = bold, head_fill
        row += 1
        for gd in rows:
            date_cell = ws.cell(row=row, column=1, value=gd.date)
            if gd.date:
                date_cell.number_format = "DD/MM/YYYY"
            else:
                date_cell.value = "Sans date"
            ws.cell(row=row, column=2, value=gd.label)
            ws.cell(row=row, column=3, value=gd.duration_hours)
            ws.cell(row=row, column=4, value=gd.billing_ref)
            row += 1
        sub = _sum_hours(rows)
        ws.cell(row=row, column=2, value="Sous-total").font = bold
        ws.cell(row=row, column=2).alignment = right
        c = ws.cell(row=row, column=3, value=sub)
        c.font, c.fill = bold, total_fill
        row += 1
        if klass.hourly_rate is not None:
            ws.cell(row=row, column=2, value="Montant (€)").alignment = right
            ws.cell(row=row, column=3, value=_money(sub * klass.hourly_rate))
            row += 1
        missing = sum(1 for gd in rows if not gd.duration_hours)
        if missing:
            ws.cell(row=row, column=1,
                    value=f"⚠ {missing} séance(s) sans durée, comptée(s) 0 h.")
            row += 1
        row += 1

    for col, width in (("A", 30), ("B", 30), ("C", 12), ("D", 30)):
        ws.column_dimensions[col].width = width
    return _xlsx_response(
        wb, f"facturation_{date_type.today().strftime('%Y-%m-%d')}.xlsx")


# --------------------------------------------------------------------------- #
# Bilan NDA : cadre C (CA) et cadre G (effectifs, heures-stagiaires)
# --------------------------------------------------------------------------- #
def nda_report(classes, start, end):
    """Chiffres du bilan pédagogique et financier sur [start, end].

    Seules comptent les séances **datées** dans la période. Un étudiant est
    compté s'il est inscrit, non neutralisé, dans une classe ayant eu au moins
    une séance sur la période ; il n'est compté qu'une fois, même inscrit dans
    plusieurs classes.

    Heures-stagiaires d'un module = heures de ses séances × étudiants actifs de
    sa classe. CA estimé = heures × taux horaire de la classe (None si le taux
    n'est pas renseigné).
    """
    modules, groups = [], {}
    all_students = set()
    for klass in _sorted_classes(classes):
        students = {e.student_id for e in klass.enrollments if e.student.active}
        level = (klass.level or "").strip() or NO_LEVEL
        for module in sorted(klass.modules, key=lambda m: m.name.lower()):
            sessions = [gd for gd in module.grade_dates
                        if gd.date and start <= gd.date <= end]
            if not sessions:
                continue
            hours = _sum_hours(sessions)
            revenue = (_money(hours * klass.hourly_rate)
                       if klass.hourly_rate is not None else None)
            row = {
                "school": klass.school.name, "level": level, "klass": klass.name,
                "module": module.name, "sessions": len(sessions),
                "missing": sum(1 for gd in sessions if not gd.duration_hours),
                "hours": hours, "students": len(students),
                "trainee_hours": round(hours * len(students), 2),
                "rate": klass.hourly_rate, "revenue": revenue,
            }
            modules.append(row)
            g = groups.setdefault((klass.school.name, level), {
                "school": klass.school.name, "level": level, "students": set(),
                "hours": 0, "trainee_hours": 0, "revenue": 0, "no_rate": 0,
            })
            g["students"] |= students
            g["hours"] += hours
            g["trainee_hours"] += row["trainee_hours"]
            if revenue is None:
                g["no_rate"] += 1
            else:
                g["revenue"] += revenue
            all_students |= students

    group_rows = sorted(groups.values(),
                        key=lambda g: (g["school"].lower(), g["level"].lower()))
    for g in group_rows:
        g["students"] = len(g["students"])
        g["hours"] = round(g["hours"], 2)
        g["trainee_hours"] = round(g["trainee_hours"], 2)
        g["revenue"] = _money(g["revenue"])
    return {
        "modules": modules, "groups": group_rows,
        "students": len(all_students),
        "hours": round(sum(m["hours"] for m in modules), 2),
        "trainee_hours": round(sum(m["trainee_hours"] for m in modules), 2),
        "revenue": _money(sum(m["revenue"] or 0 for m in modules)),
        "no_rate": sum(1 for m in modules if m["revenue"] is None),
        "missing": sum(m["missing"] for m in modules),
    }


@billing_bp.route("/facturation/nda.xlsx")
@login_required
def export_nda():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    this_year = date_type.today().year
    start = _parse_day(request.args.get("start"), date_type(this_year, 1, 1))
    end = _parse_day(request.args.get("end"), date_type(this_year, 12, 31))
    if end < start:
        start, end = end, start
    report = nda_report(visible_classes_query().all(), start, end)

    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="F1F5F9")

    def header(ws, row, titles):
        for j, title in enumerate(titles, start=1):
            c = ws.cell(row=row, column=j, value=title)
            c.font, c.fill = bold, head_fill

    period = f"du {start.strftime('%d/%m/%Y')} au {end.strftime('%d/%m/%Y')}"

    # --- Synthèse ---
    ws = wb.active
    ws.title = "Synthèse"
    ws["A1"] = f"Bilan NDA — période {period}"
    ws["A1"].font = Font(bold=True, size=13)
    lines = [
        ("Cadre G — Étudiants uniques (inscrits, non neutralisés)", report["students"]),
        ("Heures d'enseignement (séances de la période)", report["hours"]),
        ("Cadre G — Heures-stagiaires (heures × étudiants)", report["trainee_hours"]),
        ("Cadre C — CA estimé (heures × taux horaire, €)", report["revenue"]),
    ]
    for i, (label, value) in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=label).font = bold
        ws.cell(row=i, column=2, value=value)
    notes = []
    if report["no_rate"]:
        notes.append(f"⚠ {report['no_rate']} module(s) d'une classe sans taux horaire : "
                     "exclus du CA estimé.")
    if report["missing"]:
        notes.append(f"⚠ {report['missing']} séance(s) sans durée, comptée(s) 0 h.")
    notes.append("Les effectifs sont ceux du jour de l'extraction : "
                 "ré-extraire après mise à jour des inscriptions.")
    for i, note in enumerate(notes, start=8):
        ws.cell(row=i, column=1, value=note)
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 16

    # --- Cadre G : par école et par niveau ---
    ws = wb.create_sheet("École & niveau")
    header(ws, 1, ["École", "Niveau", "Étudiants uniques", "Heures d'enseignement",
                   "Heures-stagiaires", "CA estimé (€)", "Modules sans taux"])
    for r, g in enumerate(report["groups"], start=2):
        for j, key in enumerate(("school", "level", "students", "hours",
                                 "trainee_hours", "revenue", "no_rate"), start=1):
            ws.cell(row=r, column=j, value=g[key])
    total_row = len(report["groups"]) + 2
    ws.cell(row=total_row, column=1, value="Total").font = bold
    for j, key in ((3, "students"), (4, "hours"), (5, "trainee_hours"), (6, "revenue")):
        ws.cell(row=total_row, column=j, value=report[key]).font = bold
    for col, width in zip("ABCDEFG", (28, 22, 18, 22, 18, 16, 18)):
        ws.column_dimensions[col].width = width

    # --- Détail par module ---
    ws = wb.create_sheet("Par module")
    header(ws, 1, ["École", "Niveau", "Classe", "Module", "Séances", "Heures",
                   "Séances sans durée", "Étudiants actifs", "Heures-stagiaires",
                   "Taux horaire (€/h)", "CA estimé (€)"])
    for r, m in enumerate(report["modules"], start=2):
        for j, key in enumerate(("school", "level", "klass", "module", "sessions",
                                 "hours", "missing", "students", "trainee_hours",
                                 "rate", "revenue"), start=1):
            ws.cell(row=r, column=j, value=m[key])
    for col, width in zip("ABCDEFGHIJK", (24, 18, 16, 28, 10, 10, 18, 16, 18, 18, 16)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    return _xlsx_response(
        wb, f"bilan_nda_{start.strftime('%Y%m%d')}_{end.strftime('%Y%m%d')}.xlsx")
